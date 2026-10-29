"""Build editable Phase 4 project report and synopsis from measured outputs."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]


def _title(document: Document, text: str, subtitle: str = "") -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(22)
    if subtitle:
        sub = document.add_paragraph(subtitle)
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _page(document: Document, heading: str, paragraphs: list[str]) -> None:
    if len(document.paragraphs) > 1:
        document.add_page_break()
    document.add_heading(heading, level=1)
    for text in paragraphs:
        document.add_paragraph(text)


def _add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = value


def _architecture_figure(root: Path) -> Path:
    """Create a simple architecture/data-flow diagram used by both documents."""
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch

    path = root / "reports" / "architecture_diagram.png"
    figure, axis = plt.subplots(figsize=(12, 3))
    axis.axis("off")
    labels = [
        "PDF resume\n+ Job description",
        "Flask\napplication",
        "PDF text\nextraction",
        "spaCy + skill\nmatching",
        "TF-IDF +\nscoring",
        "Ranking +\nhuman review",
    ]
    for index, label in enumerate(labels):
        x = index * 1.65
        axis.add_patch(FancyBboxPatch((x, 0.35), 1.25, 0.8, boxstyle="round,pad=0.03", facecolor="#e8f1fb", edgecolor="#17365d"))
        axis.text(x + 0.625, 0.75, label, ha="center", va="center", fontsize=9)
        if index < len(labels) - 1:
            axis.annotate("", xy=(x + 1.62, 0.75), xytext=(x + 1.28, 0.75), arrowprops={"arrowstyle": "->", "color": "#17365d"})
    figure.tight_layout()
    figure.savefig(path, dpi=160)
    plt.close(figure)
    return path


def build_report(root: Path = ROOT) -> Path:
    evaluation = json.loads((root / "reports" / "evaluation.json").read_text(encoding="utf-8"))
    fairness = json.loads((root / "reports" / "fairness.json").read_text(encoding="utf-8"))
    document = Document()
    architecture = _architecture_figure(root)
    _title(document, "AI-Based Resume Screening and Candidate Shortlisting System", "HCLTech Career Shaper Programme | Project No. 34")
    document.add_paragraph("Technical project report — Phase 1 to Phase 4 prototype")
    document.add_paragraph("Prepared for academic submission. Personal and administrative fields are intentionally not populated.")

    sections = [
        ("1. Project overview", [
            "This project implements a recruiter decision-support prototype that compares a job description with uploaded PDF resumes. It extracts text, identifies explicit technical skills, computes lexical TF-IDF similarity, produces an explainable weighted score, ranks candidates, and records a separate human review decision.",
            "The application is not an autonomous hiring system. Its output is a review aid and must not be interpreted as a probability of employment or as a substitute for human judgment."
        ]),
        ("2. Problem statement and objectives", [
            "Recruiters often review many resumes against a job description under time constraints. Relevant evidence may be distributed across skills, experience, education, and descriptive text, making consistent comparison difficult.",
            "The objective is to provide a transparent local prototype that highlights matched and missing skills, exposes score components, preserves reviewer decisions separately, and supports reproducible evaluation and fairness checks."
        ]),
        ("3. Requirements and scope", [
            "In scope: eight-column dataset audit and EDA; PDF text extraction; spaCy phrase matching; domain skill synonyms; TF-IDF similarity; configurable explainable scoring; Flask upload workflow; SQLite structured persistence; review actions; CSV export; evaluation and counterfactual fairness documentation.",
            "Out of scope: DOCX resume parsing, production authentication, autonomous hiring, demographic inference, deep-learning semantic embeddings, and claims of legal compliance or real-world hiring validity."
        ]),
        ("4. Technology and literature overview", [
            "The implementation uses Python, Flask, SQLite, pdfminer.six, spaCy, scikit-learn, Pandas, NumPy, Matplotlib, Seaborn, pytest, and python-docx. TF-IDF is used as a lexical baseline because it is lightweight, inspectable, and reproducible offline.",
            "Primary technical references include the Flask documentation, pdfminer.six documentation, spaCy usage documentation, scikit-learn feature extraction documentation, and the Python sqlite3 documentation."
        ]),
        ("5. System architecture and data flow", [
            "The data flow is: Resume PDF + Job Description -> Flask Application -> PDF Text Extraction -> spaCy Preprocessing -> Skill Extraction -> TF-IDF and Cosine Similarity -> Explainable Scoring -> Candidate Ranking -> Human Review.",
            "Uploaded files are written to a temporary directory, processed, and removed. SQLite stores structured score outputs and review metadata, not raw PDFs or complete extracted resume text by default."
        ]),
        ("6. Dataset and data dictionary", [
            "The supplied dataset contains 500 rows and 8 columns: resume_id, job_role, years_exp, skills_matched, education_score, keyword_density, predicted_score, and shortlisted. It contains no nulls, duplicate rows, or duplicate resume IDs.",
            "The observed shortlist encoding is integer 0/1 although the brief describes Y/N. There are five observed job roles. The completed data dictionary documents types, meanings, ranges, units, and examples. Synthetic expansion is staged separately and does not overwrite raw data."
        ]),
        ("7. Exploratory data analysis", [
            "The supplied data has 173 shortlisted records (34.6%) and 327 non-shortlisted records (65.4%). Numeric ranges and correlations are saved in reports/audit.json.",
            "The measured correlation between predicted_score and shortlisted is -0.0085. This weak and suspicious relationship, plus uncertain label provenance, is why no supervised classifier was trained on these labels.",
        ]),
        ("8. NLP and scoring methodology", [
            "The parser normalizes Unicode whitespace and reports success, empty, encrypted, unsupported, or unreadable states. The NLP layer uses explicit configurable phrase matching and canonical synonym mapping for terms such as sklearn/scikit-learn, fast API/FastAPI, and natural language processing/NLP.",
            "The final score uses 50% text similarity, 30% required skill coverage, 15% experience match, and 5% education match. Missing components are explicitly excluded with weight renormalization. These weights are starting assumptions rather than validated optimal parameters."
        ]),
        ("9. Application design and implementation", [
            "The Flask application provides health, screening, candidate listing, detail, review, and CSV export endpoints. The dashboard accepts a job description and multiple PDFs, displays ranked results, and provides shortlist, reject, and needs-review actions.",
            "Candidate IDs are generated independently of filenames and personal names. Review decisions and notes are stored in separate fields from algorithmic scores."
        ]),
        ("10. Evaluation methodology and actual results", [
            "Six synthetic PDF cases were manually judged before execution. Relevance meant explicit coverage of the job description's required skills. This evaluation did not use dataset.csv or its labels.",
            f"Measured results: parsing success rate {evaluation['metrics']['parsing_success_rate']:.2%}; average processing time {evaluation['metrics']['average_processing_seconds']:.6f} seconds per PDF; Precision@3 {evaluation['metrics']['precision_at_3']:.4f}; Recall@3 {evaluation['metrics']['recall_at_3']:.4f}; NDCG@3 {evaluation['metrics']['ndcg_at_3']:.4f}.",
            "These results describe a small synthetic test set and do not estimate production performance."
        ]),
        ("11. Fairness and ethical assessment", [
            f"Four controlled counterfactual pairs changed names, pronouns, or gender cues while keeping qualifications constant. The measured final-score difference was non-zero for {fairness['changed_pairs']} of {fairness['pairs_tested']} pairs.",
            "The supplied dataset has no validated demographic attributes, so group-level fairness metrics were not calculated. The counterfactual result is not proof of fairness; it is a limited test of whether these cues affect this implementation.",
            "Recommendations are to minimize non-job-relevant demographic text in scoring inputs, obtain lawful and consented demographic data only for approved audits, and retain human oversight."
        ]),
        ("12. Privacy and human oversight", [
            "TemporaryDirectory is used for uploaded PDFs and cleanup occurs after processing. The database stores structured outputs, warnings, scores, and reviewer decisions, not raw PDF files or complete resume text.",
            "The project follows a minimization-oriented design, but production access control, retention schedules, deployment hardening, and a legal review are not implemented. DPDP Act 2023 considerations should be reviewed with the appropriate institution; this prototype does not independently certify legal compliance."
        ]),
        ("13. Challenges and solutions", [
            "Challenge: the supplied label relationship was not demonstrably suitable for supervised learning. Resolution: the audit measured correlations and separated matching evaluation from dataset labels.",
            "Challenge: resumes may be scanned, malformed, encrypted, or empty. Resolution: parser statuses and warnings are explicit, scores can be incomplete, and integration tests cover invalid and empty inputs.",
            "Challenge: reviewers need to disagree with automated recommendations. Resolution: review actions and notes are persisted separately from computed scores."
        ]),
        ("14. Limitations, future work, conclusion and references", [
            "Limitations include a small synthetic evaluation set, lexical rather than semantic similarity, fallback spaCy processing when the statistical model is unavailable, no DOCX support, no production authentication, and no validated demographic data.",
            "Future work should add approved authentication and retention controls, stronger semantic evaluation, larger independently reviewed datasets, accessibility testing, and a formal fairness protocol.",
            "The completed prototype demonstrates an auditable workflow from PDF upload through human review while documenting where evidence is insufficient.",
            "References: Flask documentation (https://flask.palletsprojects.com/); spaCy usage documentation (https://spacy.io/usage); scikit-learn TF-IDF documentation (https://scikit-learn.org/stable/modules/feature_extraction.html); pdfminer.six documentation (https://pdfminersix.readthedocs.io/); Python sqlite3 documentation (https://docs.python.org/3/library/sqlite3.html); Digital Personal Data Protection Act, 2023 (https://www.meity.gov.in/)."
        ]),
    ]
    for heading, paragraphs in sections:
        _page(document, heading, paragraphs)

    document.add_page_break()
    document.add_heading("Figures and captions", level=1)
    document.add_picture(str(architecture), width=Inches(6.5))
    paragraph = document.add_paragraph("Figure 0. System block diagram and Level 0 data flow.")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_dir = root / "reports" / "figures"
    for filename, caption in [
        ("shortlisted_distribution.png", "Figure 1. Distribution of shortlisted and non-shortlisted records."),
        ("predicted_score_distribution.png", "Figure 2. Distribution of supplied predicted scores."),
        ("job_role_counts.png", "Figure 3. Record counts by observed job role."),
        ("numeric_correlation_heatmap.png", "Figure 4. Correlation heatmap for numeric fields."),
    ]:
        path = figure_dir / filename
        if path.exists():
            document.add_picture(str(path), width=Inches(5.8))
            paragraph = document.add_paragraph(caption)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    output = root / "reports" / "project_report.docx"
    document.save(output)
    return output


def build_synopsis(root: Path = ROOT) -> Path:
    evaluation = json.loads((root / "reports" / "evaluation.json").read_text(encoding="utf-8"))
    document = Document()
    architecture = _architecture_figure(root)
    _title(document, "Project Synopsis — Draft", "HCLTech Career Shaper Programme | Project No. 34")
    _add_table(document, ["Field", "Value"], [
        ["Project title", "AI-Based Resume Screening and Candidate Shortlisting System"],
        ["Domain", "Digital-Data-AI / AIML"],
        ["Duration", "4 Weeks"],
        ["Learner name", "[MANUAL INPUT REQUIRED]"],
        ["Roll number", "[MANUAL INPUT REQUIRED]"],
        ["College / campus", "[MANUAL INPUT REQUIRED]"],
        ["Mentor", "[MANUAL INPUT REQUIRED]"],
    ])
    for heading, text in [
        ("1. Problem statement", "Recruiters need a consistent way to compare multiple resumes with a job description. This prototype extracts explicit evidence and presents transparent recommendations while preserving human review."),
        ("2. Proposed solution and approach", "A Flask dashboard accepts a JD and PDF resumes. pdfminer.six extracts text; spaCy phrase matching identifies configured skills; TF-IDF provides lexical similarity; a weighted scoring engine returns component explanations; SQLite records structured results and reviewer decisions."),
        ("3. Technology stack", "Python, Flask, SQLite, pdfminer.six, spaCy, scikit-learn, Pandas, NumPy, Bootstrap, JavaScript, pytest, Matplotlib and Seaborn."),
        ("4. Dataset used", "dataset.csv; 500 rows and 8 columns. Key columns: resume_id, job_role, years_exp, skills_matched, predicted_score. The raw source is preserved and the completed dictionary is in data/processed."),
        ("5. Actual evaluation results", f"Six synthetic cases were evaluated. Parsing success was {evaluation['metrics']['parsing_success_rate']:.2%}; average processing time was {evaluation['metrics']['average_processing_seconds']:.6f} seconds per PDF; Precision@3 was {evaluation['metrics']['precision_at_3']:.4f}; Recall@3 was {evaluation['metrics']['recall_at_3']:.4f}; NDCG@3 was {evaluation['metrics']['ndcg_at_3']:.4f}. These are not real-world hiring metrics."),
        ("6. Challenges and resolutions", "Uncertain supplied labels were audited rather than used blindly for classification. Invalid and empty PDFs receive explicit statuses and warnings. Human review is stored separately from automated scores."),
        ("7. Ethical declaration", "Automated scores are decision-support aids only. No demographic inference is performed. Raw PDFs are processed temporarily. DPDP Act 2023 considerations require institutional/legal review; the prototype does not certify compliance."),
        ("8. References", "Flask documentation; spaCy documentation; scikit-learn feature extraction documentation; pdfminer.six documentation; Python sqlite3 documentation."),
        ("Signature and date", "Learner signature: ____________________   Date: ________________"),
    ]:
        document.add_heading(heading, level=1)
        document.add_paragraph(text)
    document.add_heading("Architecture", level=1)
    document.add_paragraph("Resume PDF + JD -> Flask -> PDF extraction -> spaCy preprocessing -> skill extraction -> TF-IDF similarity -> explainable scoring -> ranking -> human review")
    document.add_picture(str(architecture), width=Inches(6.5))
    document.add_paragraph("Figure 1. System block diagram and Level 0 data flow.", style=None)
    output = root / "reports" / "Project_Synopsis_DRAFT_Name_RollNo.docx"
    document.save(output)
    return output


if __name__ == "__main__":
    build_report()
    build_synopsis()
